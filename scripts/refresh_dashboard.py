#!/usr/bin/env python3
"""Refresh the static weekly dashboard from a supplied leadership workbook."""

from __future__ import annotations

import argparse
import json
import re
import shutil
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook


DATA_PATTERN = re.compile(
    r"/\* DASHBOARD_DATA_START \*/.*?/\* DASHBOARD_DATA_END \*/", re.DOTALL
)
SUPPORTED_RAGS = ("Red", "Amber", "Green")
SOURCE_FILENAME = "Weekly_Leadership_Source_Workbook.xlsx"


def text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def percentage(value: object) -> int | None:
    if value in (None, ""):
        return None
    try:
        number = Decimal(str(value))
    except Exception as exc:
        raise ValueError(f"Unsupported completion value: {value!r}") from exc
    if abs(number) <= 1:
        number *= 100
    return int(number.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def rgb(color: object) -> str | None:
    if color is None or getattr(color, "type", None) != "rgb":
        return None
    value = getattr(color, "rgb", None)
    if not isinstance(value, str):
        return None
    value = value[-6:].upper()
    return f"#{value}"


def rag_from_status(status: str) -> str | None:
    for rag in SUPPORTED_RAGS:
        if re.search(rf"\b{rag}\b", status, re.IGNORECASE):
            return rag
    return None


def source_style(cell: object) -> tuple[str | None, str | None]:
    fill = None
    if getattr(cell.fill, "fill_type", None):
        fill = rgb(cell.fill.fgColor)
    font = rgb(cell.font.color)
    return fill, font


def record_from_row(sheet: object, row_number: int) -> dict[str, object] | None:
    values = [sheet.cell(row_number, column).value for column in range(1, 15)]
    if not any(value not in (None, "") for value in values):
        return None
    if values[1] in (None, ""):
        raise ValueError(f"{sheet.title}!{row_number} has data but no Project / Workstream")

    explicit_rag = text(values[3]).title()
    if explicit_rag in SUPPORTED_RAGS:
        status = text(values[2])
        rag = explicit_rag
        style_cell = sheet.cell(row_number, 4)
        start, finish, completion = values[4], values[5], values[6]
        objective, progress, next_step = values[7], values[8], values[9]
        owner, risks, attention, source = values[10], values[11], values[12], values[13]
    else:
        status = text(values[2])
        rag = rag_from_status(status)
        style_cell = sheet.cell(row_number, 3) if rag else None
        start, finish, completion = values[3], values[4], values[5]
        objective, progress, next_step = values[6], values[7], values[8]
        owner, risks, attention, source = values[9], values[10], values[11], values[12]

    rag_fill, rag_font = source_style(style_cell) if style_cell else (None, None)
    return {
        "id": text(values[0]),
        "project": text(values[1]),
        "team": sheet.title,
        "status": status,
        "rag": rag,
        "ragFill": rag_fill,
        "ragFont": rag_font,
        "start": text(start),
        "finish": text(finish),
        "completion": percentage(completion),
        "objective": text(objective),
        "progress": text(progress),
        "next": text(next_step),
        "owner": text(owner),
        "risks": text(risks),
        "attention": text(attention),
        "source": text(source),
        "sourceRow": row_number,
    }


def build_records(workbook_path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(workbook_path, data_only=False, read_only=False)
    records: list[dict[str, object]] = []
    for sheet in workbook.worksheets:
        for row_number in range(2, sheet.max_row + 1):
            record = record_from_row(sheet, row_number)
            if record:
                record["index"] = len(records)
                records.append(record)
    return records


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--report-date", required=True, help="ISO date for the weekly report")
    args = parser.parse_args()

    workbook_path = args.workbook.resolve()
    if not workbook_path.is_file():
        raise SystemExit(f"Workbook not found: {workbook_path}")
    try:
        datetime.strptime(args.report_date, "%Y-%m-%d")
    except ValueError as exc:
        raise SystemExit("--report-date must use YYYY-MM-DD") from exc

    project_dir = Path(__file__).resolve().parent.parent
    index_path = project_dir / "index.html"
    records = build_records(workbook_path)
    if not records:
        raise SystemExit("No valid workbook records found")

    payload = {
        "reportDate": args.report_date,
        "sourceFile": SOURCE_FILENAME,
        "dashboardFile": "IT_Leadership_Dashboard_Weekly.html",
        "records": records,
    }
    html = index_path.read_text(encoding="utf-8")
    replacement = f"/* DASHBOARD_DATA_START */{json.dumps(payload, ensure_ascii=True, separators=(',', ':'))}/* DASHBOARD_DATA_END */"
    updated, count = DATA_PATTERN.subn(lambda _match: replacement, html, count=1)
    if count != 1:
        raise SystemExit("Dashboard data marker was not found exactly once")
    index_path.write_text(updated, encoding="utf-8")
    shutil.copyfile(workbook_path, project_dir / SOURCE_FILENAME)
    print(f"Updated {index_path} with {len(records)} records")
    print(f"Copied source workbook to {project_dir / SOURCE_FILENAME}")


if __name__ == "__main__":
    main()
